from io import BytesIO
import re

from docx import Document
from docx.shared import Pt
from fpdf import FPDF


def format_docx(text, doc_type, terms=""):
    doc = Document()

    title = doc.add_heading(doc_type, 0)
    title.alignment = 1

    for line in text.splitlines():
        if line.strip():
            paragraph = doc.add_paragraph(line.strip())

            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)

    if terms:
        doc.add_heading("Terms Summary", 1)

        for i, term in enumerate(terms.split(";"), 1):
            if term.strip():
                doc.add_paragraph(
                    f"{i}. {term.strip()}"
                )

    output = BytesIO()
    doc.save(output)

    return output.getvalue()


def clean_text(text):
    text = str(text)

    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2026": "...",
        "\u00a0": " ",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    return text.encode(
        "latin-1",
        "replace"
    ).decode("latin-1")


def safe_lines(text, max_chars=70):
    """
    Break long text into safe pieces for FPDF.
    """
    result = []

    for line in text.splitlines():

        line = clean_text(line.strip())

        if not line:
            result.append("")
            continue

        # Break long repeated characters
        line = re.sub(
            r"([_=+\-])\1{10,}",
            lambda m: m.group(1) * 10,
            line
        )

        # Split long lines into smaller pieces
        while len(line) > max_chars:
            result.append(line[:max_chars])
            line = line[max_chars:]

        if line:
            result.append(line)

    return result


def format_pdf(text, doc_type, terms=""):
    pdf = FPDF()

    pdf.set_auto_page_break(
        auto=True,
        margin=20
    )

    pdf.set_margins(
        20,
        20,
        20
    )

    pdf.add_page()

    # -------------------------
    # TITLE
    # -------------------------

    pdf.set_font(
        "Helvetica",
        "B",
        16
    )

    title = clean_text(doc_type)

    pdf.set_x(pdf.l_margin)

    pdf.multi_cell(
        pdf.epw,
        10,
        title,
        align="C"
    )

    pdf.ln(8)

    # -------------------------
    # DOCUMENT CONTENT
    # -------------------------

    pdf.set_font(
        "Helvetica",
        "",
        11
    )

    for line in safe_lines(text):

        if not line:
            pdf.ln(4)
            continue

        pdf.set_x(pdf.l_margin)

        pdf.multi_cell(
            pdf.epw,
            6,
            line,
            align="L"
        )

    # -------------------------
    # TERMS SUMMARY
    # -------------------------

    if terms:

        pdf.ln(5)

        pdf.set_font(
            "Helvetica",
            "B",
            12
        )

        pdf.set_x(pdf.l_margin)

        pdf.multi_cell(
            pdf.epw,
            8,
            "Terms Summary",
            align="L"
        )

        pdf.set_font(
            "Helvetica",
            "",
            10
        )

        for i, term in enumerate(
            terms.split(";"),
            1
        ):

            term = clean_text(term.strip())

            if term:

                pdf.set_x(pdf.l_margin)

                pdf.multi_cell(
                    pdf.epw,
                    6,
                    f"{i}. {term}",
                    align="L"
                )

    return bytes(pdf.output())
